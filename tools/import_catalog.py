"""Importación cotejada por fila: las anclas contienen fotos alternativas ajenas al producto.
Conserva las rutas públicas usadas por JSONBin. No modifica el catálogo remoto.
"""
import openpyxl,json,re,io,hashlib,argparse
from pathlib import Path
from PIL import Image
root=Path(__file__).resolve().parent.parent
parser=argparse.ArgumentParser();parser.add_argument('source',nargs='?',default=r'C:\Users\USUARIO\Downloads\CATALOGO (1).xlsx');args=parser.parse_args()
source=Path(args.source);source_hash=hashlib.sha256(source.read_bytes()).hexdigest()
report_file=root/'data/catalog.review.json'
if report_file.exists() and json.loads(report_file.read_text(encoding='utf-8'))['source_sha256']!=source_hash:
    raise ValueError('El Excel cambió. Debes revisar visualmente las fotos antes de reutilizar este mapa.')
s=openpyxl.load_workbook(source,data_only=True).active
# Índices de imágenes contrastados visualmente con el nombre de cada fila.
selected={2:1,3:3,4:8,5:11,6:13,7:15,8:18,9:20,10:28,11:32,12:36,13:38,14:39,15:41,16:42,17:43,18:46,19:47,20:45,21:44}
# Destinos originales: JSONBin ya usa estas rutas. Cambiar la foto, no el enlace.
primary={2:1,3:3,4:4,5:11,6:12,7:15,8:16,9:20,10:24,11:30,12:34,13:38,14:39,15:40,16:42,17:43,18:46,19:47,20:45,21:44}
rows=[r for r in range(2,s.max_row+1) if s.cell(r,3).value]
assert rows==list(selected),'Cambió la estructura del archivo. Requiere nuevo cotejo.'
(root/'assets/catalogo').mkdir(exist_ok=True)
products=[];report=[];category=''
for row in rows:
    cat,_,name,description,price,old=[s.cell(row,c).value for c in range(1,7)]
    category=str(cat).strip() if cat else category
    image=s._images[selected[row]-1]
    assert image.anchor._from.row+1==row,f'Ancla inesperada en fila {row}'
    raw=image._data();photo=Image.open(io.BytesIO(raw)).convert('RGB');photo.thumbnail((1200,1200))
    relative=f'assets/catalogo/producto-{row-1}-{primary[row]}.webp';photo.save(root/relative,'WEBP',quality=86)
    description=str(description or '').replace('Haz la descripción más vendedoraAñade una llamada a la acción','').strip()
    if row==17:description=description.replace(' Se entrega certificado al concluir el curso','')
    p=dict(id=1000+row,nombre=re.sub(r'\s+',' ',name).strip(),categoria=category,precio=float(price),precioAnterior=float(old) if old else None,imagen=relative,imagenes=[relative],descripcion=description,disponible='SI',destacado='SI' if old and old>price else 'NO')
    products.append(p);report.append(dict(fila=row,id=p['id'],nombre=p['nombre'],categoria=category,precio=p['precio'],imagen_excel=selected[row],imagen_publicada=relative,imagen_excel_sha256=hashlib.sha256(raw).hexdigest()))
config=dict(negocio='Winni Store',whatsapp='51914663723',moneda='S/',catalogoDemo=False)
(root/'config.js').write_text('// Catálogo cotejado visualmente por fila. Una imagen por producto.\nconst CONFIG = '+json.dumps(config,ensure_ascii=False)+';\nconst PRODUCTOS = '+json.dumps(products,ensure_ascii=False,indent=2)+';\n',encoding='utf-8')
(root/'data/catalog.seed.json').write_text(json.dumps(dict(products=products),ensure_ascii=False,indent=2),encoding='utf-8')
report_file.write_text(json.dumps(dict(source_sha256=source_hash,hoja=s.title,productos=report,observaciones=['Fila 9: categoría vacía, se conserva CALZADO Y ZAPATILLAS de su sección.','D13: omitida una instrucción editorial del texto público.','D17: omitida una frase de certificado ajena a la prenda.']),ensure_ascii=False,indent=2),encoding='utf-8')
print(f'{len(products)} productos cotejados, rutas estables conservadas. Sin escrituras en JSONBin.')
